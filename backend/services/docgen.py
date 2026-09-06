"""Renders the structured JSON produced by the doc agents into real Office files.

BRD/FRD are narrative requirement documents -> real .docx (python-docx).
Data Dictionary/Integration Mapping are inherently tabular -> real .xlsx (openpyxl).

Deterministic, template-driven — no LLM call happens here. This is purely a
formatting layer over content the agents already produced.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ---- Brand palette (matches the frontend's design tokens) ----
INK = RGBColor(0x16, 0x22, 0x23)
MUTED = RGBColor(0x70, 0x80, 0x80)
TEAL = RGBColor(0x0E, 0x6B, 0x63)
AMBER = RGBColor(0xA6, 0x72, 0x0C)
VIOLET = RGBColor(0x5B, 0x4A, 0x96)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

TEAL_HEX = "0E6B63"
TEAL_SOFT_HEX = "E7F4F1"
AMBER_HEX = "A6720C"
AMBER_SOFT_HEX = "FFF4D9"
VIOLET_HEX = "5B4A96"
RED_HEX = "B4232B"
RED_SOFT_HEX = "FCE8E8"
LINE_HEX = "DCE5E4"
INK_HEX = "162223"


# =====================================================================  DOCX

def _cell_shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
    tcPr.append(shd)


def _set_cell_text(cell, text: str, bold=False, color=None, size=10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = color


def _new_doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK
    for section in doc.sections:
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
    return doc


def _cover(doc: Document, kicker: str, title: str, meta: Dict[str, Any]) -> None:
    p = doc.add_paragraph()
    r = p.add_run("SemantIQ")
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = TEAL

    p = doc.add_paragraph()
    r = p.add_run(kicker)
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = AMBER
    p.space_before = Pt(14)

    p = doc.add_paragraph()
    r = p.add_run(title)
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = INK
    p.space_before = Pt(4)
    p.space_after = Pt(10)

    p = doc.add_paragraph()
    r = p.add_run(
        f"Prepared from {meta.get('source_label', 'the connected source model')} and business "
        f"context captured on {meta.get('generated_date', '')}."
    )
    r.font.size = Pt(10.5)
    r.font.italic = True
    r.font.color.rgb = MUTED
    p.space_after = Pt(4)

    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    labels = ["Version", "Status", "Prepared for", "Platform"]
    values = ["1.0", "Draft — for review", meta.get("audience", "Stakeholders"), meta.get("platform", "—")]
    for i, (lab, val) in enumerate(zip(labels, values)):
        _set_cell_text(table.cell(0, i), lab, bold=True, color=MUTED, size=9)
        _set_cell_text(table.cell(1, i), val, size=10)
        _cell_shade(table.cell(0, i), TEAL_SOFT_HEX)
    doc.add_paragraph().space_after = Pt(6)


def _heading(doc: Document, text: str, level=1, color=TEAL) -> None:
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    r.font.color.rgb = color
    r.font.name = "Calibri"
    if level == 1:
        r.font.size = Pt(16)
        p.space_before = Pt(18)
        p.space_after = Pt(6)
    else:
        r.font.size = Pt(13)
        p.space_before = Pt(12)
        p.space_after = Pt(4)


def _para(doc: Document, text: str, italic=False, color=INK, size=10.5) -> None:
    if not text:
        return
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.color.rgb = color
    p.space_after = Pt(6)


def _bullets(doc: Document, items: List[str]) -> None:
    for item in items or []:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(str(item))
        r.font.size = Pt(10.5)
        r.font.color.rgb = INK


def _table(doc: Document, headers: List[str], rows: List[List[str]], widths=None, header_color=TEAL_HEX) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        _set_cell_text(table.cell(0, i), h, bold=True, color=WHITE, size=9.5)
        _cell_shade(table.cell(0, i), header_color)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            _set_cell_text(cells[i], val, size=9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().space_after = Pt(4)


def _priority_label(p: str) -> str:
    return {"Must": "MUST HAVE", "Should": "SHOULD HAVE", "Could": "COULD HAVE"}.get(p, p or "—")


def build_brd_docx(doc_json: Dict[str, Any], meta: Dict[str, Any]) -> bytes:
    doc = _new_doc()
    _cover(doc, doc_json.get("kicker", "BRD"), doc_json.get("title", "Business Requirements Document"), meta)

    _heading(doc, "01 · Executive Summary")
    _para(doc, doc_json.get("executive_summary", ""))

    _heading(doc, "02 · Business Background")
    _para(doc, doc_json.get("business_background", ""))

    _heading(doc, "03 · Business Objectives")
    objs = doc_json.get("objectives", [])
    if objs:
        _table(doc, ["ID", "Objective", "Success Metric"],
               [[o.get("id", ""), o.get("statement", ""), o.get("success_metric", "")] for o in objs],
               widths=[0.7, 3.5, 2.5])

    _heading(doc, "04 · Stakeholders")
    sh = doc_json.get("stakeholders", [])
    if sh:
        _table(doc, ["Role", "Responsibility", "Interest"],
               [[s.get("role", ""), s.get("responsibility", ""), s.get("interest", "")] for s in sh],
               widths=[1.8, 2.7, 2.2])

    _heading(doc, "05 · Current State vs. Future State")
    _heading(doc, "Current State", level=2, color=MUTED)
    _para(doc, doc_json.get("current_state", ""))
    _heading(doc, "Future State", level=2, color=TEAL)
    _para(doc, doc_json.get("future_state", ""))

    _heading(doc, "06 · Scope")
    _heading(doc, "In Scope", level=2, color=TEAL)
    _bullets(doc, doc_json.get("in_scope", []))
    _heading(doc, "Out of Scope", level=2, color=AMBER)
    _bullets(doc, doc_json.get("out_of_scope", []))

    _heading(doc, "07 · Business Requirements")
    reqs = doc_json.get("requirements", [])
    themes = list(dict.fromkeys(r.get("theme", "General") for r in reqs)) or ["General"]
    for theme in themes:
        _heading(doc, theme, level=2, color=VIOLET)
        theme_reqs = [r for r in reqs if r.get("theme", "General") == theme]
        _table(doc, ["ID", "Requirement", "Priority", "Rationale"],
               [[r.get("id", ""), r.get("description", ""), _priority_label(r.get("priority", "")), r.get("rationale", "")] for r in theme_reqs],
               widths=[0.7, 2.6, 1.0, 2.4])

    _heading(doc, "08 · Assumptions & Constraints")
    _heading(doc, "Assumptions", level=2, color=MUTED)
    _bullets(doc, doc_json.get("assumptions", []))
    _heading(doc, "Constraints", level=2, color=MUTED)
    _bullets(doc, doc_json.get("constraints", []))

    _heading(doc, "09 · Risks & Mitigations")
    risks = doc_json.get("risks", [])
    if risks:
        _table(doc, ["Risk", "Impact", "Mitigation"],
               [[r.get("risk", ""), r.get("impact", ""), r.get("mitigation", "")] for r in risks],
               widths=[2.5, 1.0, 3.2])

    _heading(doc, "10 · Success Criteria")
    _bullets(doc, doc_json.get("success_criteria", []))

    _heading(doc, "11 · Glossary")
    gl = doc_json.get("glossary", [])
    if gl:
        _table(doc, ["Term", "Definition"], [[g.get("term", ""), g.get("definition", "")] for g in gl], widths=[1.8, 4.9])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_frd_docx(doc_json: Dict[str, Any], meta: Dict[str, Any]) -> bytes:
    doc = _new_doc()
    _cover(doc, doc_json.get("kicker", "FRD"), doc_json.get("title", "Functional Requirements Document"), meta)

    _heading(doc, "01 · Introduction")
    _para(doc, doc_json.get("introduction", ""))

    _heading(doc, "02 · System Overview")
    _para(doc, doc_json.get("system_overview", ""))

    _heading(doc, "03 · Functional Modules")
    for mi, module in enumerate(doc_json.get("modules", []), start=1):
        _heading(doc, f"Module {mi} · {module.get('name', '')}", level=2, color=VIOLET)
        _para(doc, module.get("description", ""), color=MUTED)
        for story in module.get("stories", []):
            p = doc.add_paragraph()
            r = p.add_run(f"{story.get('id', '')}  ")
            r.font.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = TEAL
            r2 = p.add_run(story.get("story", ""))
            r2.font.size = Pt(10.5)
            r3 = p.add_run(f"   [{_priority_label(story.get('priority', ''))}]")
            r3.font.size = Pt(8.5)
            r3.font.bold = True
            r3.font.color.rgb = AMBER
            p.space_after = Pt(2)
            for crit in story.get("acceptance_criteria", []):
                cp = doc.add_paragraph(style="List Bullet 2")
                cr = cp.add_run(str(crit))
                cr.font.size = Pt(9.5)
                cr.font.color.rgb = MUTED
            doc.add_paragraph().space_after = Pt(2)

    _heading(doc, "04 · Non-Functional Requirements")
    nfr = doc_json.get("non_functional_requirements", [])
    if nfr:
        _table(doc, ["Category", "Requirement"], [[n.get("category", ""), n.get("requirement", "")] for n in nfr], widths=[1.5, 5.2])

    _heading(doc, "05 · Data Requirements")
    dr = doc_json.get("data_requirements", [])
    if dr:
        _table(doc, ["Entity", "Fields", "Notes"],
               [[d.get("entity", ""), ", ".join(d.get("fields", [])), d.get("notes", "")] for d in dr],
               widths=[1.6, 2.8, 2.3])

    _heading(doc, "06 · Traceability Matrix")
    tr = doc_json.get("traceability", [])
    if tr:
        _table(doc, ["Business Requirement", "Functional Requirement"],
               [[t.get("business_requirement", ""), t.get("functional_requirement", "")] for t in tr],
               widths=[3.0, 3.0])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# =====================================================================  XLSX

_HEADER_FILL = PatternFill(start_color=TEAL_HEX, end_color=TEAL_HEX, fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_TITLE_FONT = Font(color=INK_HEX, bold=True, size=16)
_SUB_FONT = Font(color="708080", italic=True, size=10)
_THIN = Side(style="thin", color=LINE_HEX)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_STATUS_FILL = {
    "Matched": PatternFill(start_color=TEAL_SOFT_HEX, end_color=TEAL_SOFT_HEX, fill_type="solid"),
    "Renamed": PatternFill(start_color=AMBER_SOFT_HEX, end_color=AMBER_SOFT_HEX, fill_type="solid"),
    "Missing": PatternFill(start_color=RED_SOFT_HEX, end_color=RED_SOFT_HEX, fill_type="solid"),
}
_STATUS_FONT = {
    "Matched": Font(color=TEAL_HEX, bold=True, size=9.5),
    "Renamed": Font(color=AMBER_HEX, bold=True, size=9.5),
    "Missing": Font(color=RED_HEX, bold=True, size=9.5),
}


def _sheet_title(ws: Worksheet, title: str, subtitle: str, cols: int) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(cols, 1))
    c = ws.cell(row=1, column=1, value=title)
    c.font = _TITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(cols, 1))
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.font = _SUB_FONT
    return 4  # first header row


def _write_table(ws: Worksheet, start_row: int, headers: List[str], rows: List[List[Any]], widths: List[int]) -> None:
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=j, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        cell.border = _BORDER
    for i, row in enumerate(rows, start=start_row + 1):
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if headers[j - 1] == "Status" and val in _STATUS_FILL:
                cell.fill = _STATUS_FILL[val]
                cell.font = _STATUS_FONT[val]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def build_dictionary_xlsx(doc_json: Dict[str, Any], meta: Dict[str, Any]) -> bytes:
    wb = Workbook()
    summary = doc_json.get("summary", {})

    ws = wb.active
    ws.title = "Summary"
    row = _sheet_title(ws, "Semantic Data Dictionary", f"Prepared from {meta.get('source_label', '')} on {meta.get('generated_date', '')}", 2)
    _write_table(ws, row, ["Metric", "Count"], [
        ["Tables", summary.get("tables", 0)],
        ["Columns", summary.get("columns", 0)],
        ["Measures", summary.get("measures", 0)],
        ["Relationships", summary.get("relationships", 0)],
    ], [24, 12])

    ws2 = wb.create_sheet("Tables & Columns")
    row2 = _sheet_title(ws2, "Tables & Columns", "One row per column, grouped by table", 5)
    rows = []
    for t in doc_json.get("tables", []):
        for c in t.get("columns", []):
            rows.append([t.get("name", ""), t.get("purpose", ""), c.get("name", ""), c.get("type", ""), "Yes" if c.get("key") else "", c.get("description", "")])
    _write_table(ws2, row2, ["Table", "Table Purpose", "Column", "Type", "Key", "Column Description"], rows, [18, 34, 20, 14, 6, 34])

    ws3 = wb.create_sheet("Measures")
    row3 = _sheet_title(ws3, "Measures & KPIs", "Every DAX measure defined in the source model", 3)
    m_rows = [[m.get("name", ""), m.get("expression", ""), m.get("description", "")] for m in doc_json.get("measures", [])]
    _write_table(ws3, row3, ["Measure", "DAX Expression", "Business Description"], m_rows, [22, 40, 40])

    ws4 = wb.create_sheet("Relationships")
    row4 = _sheet_title(ws4, "Relationships", "Joins between tables in the source model", 3)
    r_rows = [[r.get("from", ""), r.get("to", ""), r.get("cardinality", "")] for r in doc_json.get("relationships", [])]
    _write_table(ws4, row4, ["From", "To", "Cardinality"], r_rows, [26, 26, 18])

    ws5 = wb.create_sheet("Glossary")
    row5 = _sheet_title(ws5, "Business Glossary", "Terms used across this documentation set", 2)
    g_rows = [[g.get("term", ""), g.get("definition", "")] for g in doc_json.get("glossary", [])]
    _write_table(ws5, row5, ["Term", "Definition"], g_rows, [22, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_mapping_xlsx(doc_json: Dict[str, Any], meta: Dict[str, Any]) -> bytes:
    wb = Workbook()
    summary = doc_json.get("summary", {})

    ws = wb.active
    ws.title = "Summary"
    row = _sheet_title(ws, "Integration Mapping Report", f"Power BI → {doc_json.get('platform', '')} · estimate pending live catalog comparison", 2)
    _write_table(ws, row, ["Metric", "Value"], [
        ["Target platform", doc_json.get("platform", "")],
        ["Mapped (Matched + Renamed)", f"{doc_json.get('mapped_pct', 0)}%"],
        ["Total fields", summary.get("total", 0)],
        ["Matched", summary.get("matched", 0)],
        ["Renamed", summary.get("renamed", 0)],
        ["Missing / needs review", summary.get("missing", 0)],
    ], [30, 20])

    ws2 = wb.create_sheet("Table Mapping")
    row2 = _sheet_title(ws2, "Table Mapping", "Source table → physical target table", 3)
    t_rows = [[t.get("source_table", ""), t.get("target_table", ""), t.get("type", "")] for t in doc_json.get("table_mappings", [])]
    _write_table(ws2, row2, ["Source Table", "Target Table", "Type"], t_rows, [22, 26, 14])

    ws3 = wb.create_sheet("Field Mapping")
    row3 = _sheet_title(ws3, "Field Mapping", "Every source column, one row each", 7)
    f_rows = [[
        f.get("source_table", ""), f.get("source_column", ""), f.get("source_type", ""),
        f.get("target_table", ""), f.get("target_column", ""), f.get("status", ""), f.get("note", ""),
    ] for f in doc_json.get("field_mappings", [])]
    _write_table(ws3, row3, ["Source Table", "Source Column", "Source Type", "Target Table", "Target Column", "Status", "Note"], f_rows, [18, 18, 12, 22, 22, 12, 34])

    ws4 = wb.create_sheet("Relationships")
    row4 = _sheet_title(ws4, "Relationship Mapping", "Source joins → target join conditions", 3)
    rel_rows = [[r.get("source", ""), r.get("target_join", ""), r.get("cardinality", "")] for r in doc_json.get("relationship_mappings", [])]
    _write_table(ws4, row4, ["Source", "Target Join", "Cardinality"], rel_rows, [30, 40, 16])

    ws5 = wb.create_sheet("Measures")
    row5 = _sheet_title(ws5, "Measure Mapping", "DAX measures to re-derive on the target platform", 3)
    meas_rows = [[m.get("measure", ""), m.get("dax", ""), m.get("note", "")] for m in doc_json.get("measure_mappings", [])]
    _write_table(ws5, row5, ["Measure", "Source DAX", "Target Notes"], meas_rows, [22, 36, 40])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DOC_BUILDERS = {
    "brd": ("docx", build_brd_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "frd": ("docx", build_frd_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "dictionary": ("xlsx", build_dictionary_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "mapping": ("xlsx", build_mapping_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}
